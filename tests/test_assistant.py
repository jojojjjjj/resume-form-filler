"""Synthetic fixtures only. No personal resume or live recruitment site is read."""
import contextlib
import importlib.util
import io
import json
from pathlib import Path
import subprocess
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch

REPO = Path(__file__).resolve().parents[1]
SCRIPTS = REPO / 'skills/resume-form-filler/scripts'
sys.path.insert(0, str(SCRIPTS))
import context
import ingest
import workflow as w
import browser_ops as b

def module(path, name):
    spec = importlib.util.spec_from_file_location(name, path)
    value = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(value)
    return value

installer = module(REPO / 'tools/install.py', 'installer')

def make_pdf(path, text='Alex Example | alex@example.com | Example University'):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
    stream = DecodedStreamObject()
    stream.set_data(('BT /F1 12 Tf 50 700 Td (' + text + ') Tj ET').encode('ascii'))
    page[NameObject('/Contents')] = writer._add_object(stream)
    with path.open('wb') as handle:
        writer.write(handle)

class ProjectCase(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='rff-test-')
        self.root = Path(self.temp.name).resolve()
        state = self.root / '.resume-form-filler'
        profile = state / 'profile.md'
        self.patches = [patch.multiple(context, ROOT=self.root, STATE=state, PROFILE=profile),
                        patch.multiple(ingest, ROOT=self.root, STATE=state, PROFILE=profile),
                        patch.multiple(w, ROOT=self.root, STATE=state, PROFILE=profile, SITES=state/'sites', RUNS=state/'runs')]
        for item in self.patches:
            item.start()
        context.init_state()
        self.tex = self.root / 'resume.tex'
        self.tex.write_text(r'\documentclass{article}\begin{document}Alex Example | alex@example.com\section{Education}Example University\end{document}', encoding='utf-8')

    def tearDown(self):
        for item in reversed(self.patches):
            item.stop()
        self.temp.cleanup()

    def imported(self):
        result = ingest.import_resume(self.tex)
        draft = Path(result['draft'])
        data = json.loads(draft.read_text())
        data['facts']['basic'].update(name='Alex Example', email='alex@example.com')
        data['facts']['education'] = [{'school': 'Example University', 'gpa': None}]
        data['evidence'] = {'basic.name': 'Alex Example', 'education.0.school': 'Example University'}
        ingest.write_json(draft, data)
        return draft, data

    def ready(self):
        draft, data = self.imported()
        ingest.review_profile(draft, 'Checked synthetic source')
        return data

    def test_pdf_text_import(self):
        path = self.root / 'resume.pdf'
        make_pdf(path)
        text, sources, warnings = ingest.extract(path)
        self.assertIn('Alex Example', text)
        self.assertEqual(sources, [path])

    def test_scanned_pdf_blocks(self):
        from pypdf import PdfWriter
        path = self.root / 'scan.pdf'
        writer = PdfWriter(); writer.add_blank_page(width=612, height=792)
        with path.open('wb') as handle: writer.write(handle)
        with self.assertRaisesRegex(ValueError, 'OCR_REQUIRED'): ingest.extract(path)

    def test_encrypted_pdf_blocks(self):
        from pypdf import PdfWriter
        writer = PdfWriter(); writer.add_blank_page(width=612, height=792); writer.encrypt('test-password')
        path = self.root / 'locked.pdf'
        with path.open('wb') as handle: writer.write(handle)
        with self.assertRaisesRegex(ValueError, 'Password-protected'): ingest.extract(path)

    def test_word_paragraph_table_header(self):
        from docx import Document
        document = Document()
        document.add_paragraph('Alex Example')
        document.add_table(rows=1, cols=1).cell(0, 0).text = 'Example University'
        document.add_paragraph('After table')
        document.sections[0].header.paragraphs[0].text = 'Header example'
        path = self.root / 'resume.docx'; document.save(path)
        text, _, _ = ingest.extract(path)
        self.assertLess(text.index('Example University'), text.index('After table'))
        self.assertIn('Header example', text)

    def test_tex_include_and_fingerprints(self):
        self.tex.write_text(r'\begin{document}Alex Example\input{education}\end{document}')
        child = self.root / 'education.tex'; child.write_text('Example University')
        draft, _ = self.imported()
        ingest.review_profile(draft, 'Checked synthetic include')
        self.assertEqual(len(w.current_profile()['sources']), 2)
        child.write_text('Changed university')
        with self.assertRaisesRegex(ValueError, 'CV_SOURCE_CHANGED'): w.current_profile()

    def test_tex_no_execution_or_external_include(self):
        outside = self.root.parent / 'outside-example.tex'
        self.tex.write_text(r'\input{../outside-example}')
        with self.assertRaisesRegex(ValueError, 'inside'): ingest.extract(self.tex)

    def test_tex_cycles(self):
        self.tex.write_text(r'\input{resume}')
        with self.assertRaisesRegex(ValueError, 'Cyclic'): ingest.extract(self.tex)

    def test_no_profile_cannot_fill(self):
        with self.assertRaisesRegex(ValueError, 'PROFILE_REQUIRED'): w.current_profile()

    def test_import_does_not_enable_filling(self):
        draft, _ = self.imported()
        self.assertFalse(w.PROFILE.exists())
        self.assertTrue(draft.exists())
        with self.assertRaisesRegex(ValueError, 'PROFILE_REQUIRED'): w.current_profile()

    def test_review_persists_reusable_profile(self):
        self.ready()
        self.assertEqual(w.current_profile()['facts']['basic']['name'], 'Alex Example')
        self.assertEqual(w.current_profile()['status'], 'ready')

    def test_unknown_name_cannot_be_ready(self):
        draft, data = self.imported(); data['facts']['basic']['name'] = None; ingest.write_json(draft, data)
        with self.assertRaisesRegex(ValueError, 'name'): ingest.review_profile(draft, 'Checked')

    def test_source_manifest_cannot_be_replaced(self):
        draft, data = self.imported(); data['sources'] = []; ingest.write_json(draft, data)
        with self.assertRaisesRegex(ValueError, 'manifest'): ingest.review_profile(draft, 'Checked')

    def test_changed_resume_requires_reimport(self):
        draft, _ = self.imported(); self.tex.write_text('changed')
        with self.assertRaisesRegex(ValueError, 'CV_SOURCE_CHANGED'): ingest.review_profile(draft, 'Checked')

    def test_replacement_backs_up_profile(self):
        self.ready(); original = w.PROFILE.read_bytes(); draft, _ = self.imported()
        with self.assertRaisesRegex(ValueError, 'PROFILE_EXISTS'): ingest.review_profile(draft, 'Checked')
        ingest.review_profile(draft, 'Checked', replace=True)
        backups = list((ingest.STATE/'profile-backups').glob('*.md'))
        self.assertEqual(backups[0].read_bytes(), original)

    def test_password_profile_is_rejected(self):
        draft, data = self.imported(); data['facts']['basic']['password'] = 'synthetic-password'; ingest.write_json(draft, data)
        with self.assertRaisesRegex(ValueError, 'Sensitive'): ingest.review_profile(draft, 'Checked')

    def test_pending_fact_cannot_resolve(self):
        self.ready(); data = w.current_profile(); data['pending_confirmations']['basic.email'] = 'ambiguous'
        with self.assertRaisesRegex(ValueError, 'Pending'): w.resolve('profile:basic.email', data, {})

    def test_identity_not_in_default_profile(self):
        self.ready(); data = w.current_profile()
        self.assertFalse(data['workflow_preferences']['identity_storage_and_fill_authorized'])
        self.assertNotIn('id_card', data['facts']['basic'])

    def test_new_site_overrides_not_guessed(self):
        self.ready()
        with self.assertRaisesRegex(ValueError, 'Unconfirmed site'): w.resolve('site:degree', w.current_profile(), {})

    def test_install_and_upgrade_preserve_state(self):
        project = self.root / 'target'; result = installer.install(project)
        personal = project / '.resume-form-filler/profile.md'; personal.write_text('local synthetic state')
        self.assertTrue(Path(result['skill'], 'scripts/ingest.py').exists())
        with self.assertRaisesRegex(ValueError, 'already exists'): installer.install(project)
        upgraded = installer.install(project, replace=True)
        self.assertTrue(Path(upgraded['backup']).is_dir())
        self.assertEqual(personal.read_text(), 'local synthetic state')
        self.assertIn('/.resume-form-filler/', (project/'.gitignore').read_text())

    def test_installed_project_path_is_discovered(self):
        project = self.root / 'target'; result = installer.install(project)
        env = dict(__import__('os').environ); env.pop('RF_PROJECT_ROOT', None)
        path = Path(result['skill']) / 'scripts'
        output = subprocess.check_output([sys.executable, '-c', 'import sys;sys.path.insert(0,sys.argv[1]);from context import ROOT;print(ROOT)', str(path)], env=env, text=True)
        self.assertEqual(Path(output.strip()), project)

    def test_installer_preserves_unmanaged_entrypoint(self):
        project = self.root / 'target'; (project/'tools').mkdir(parents=True)
        (project/'tools/resume-fill.py').write_text('personal existing tool')
        with self.assertRaisesRegex(ValueError, 'unmanaged'): installer.install(project)
        self.assertEqual((project/'tools/resume-fill.py').read_text(), 'personal existing tool')

    def test_start_rejects_credentials_before_browser(self):
        args = SimpleNamespace(url='https://example.com/edit?token=example', target_id=None)
        with patch.object(w, 'harness', side_effect=AssertionError('Browser should not be contacted')):
            with self.assertRaisesRegex(ValueError, 'tokens'): w.start(args)

    def test_stale_plan_binding_is_rejected(self):
        self.ready()
        run = ingest.STATE/'runs/test'; run.mkdir()
        session = {'interface_id': 'local', 'url': 'https://example.com/edit', 'target_id': 'fixture', 'profile_sha256': w.digest(w.PROFILE), 'site_sha256': None, 'site_path': None}
        w.save_json(run/'session.json', session)
        w.save_json(run/'plan.json', {**session, 'profile_sha256': 'old', 'fields': []})
        args = SimpleNamespace(session=str(run), plan=str(run/'plan.json'), cdp_url=None, cdp_ws=None)
        with self.assertRaisesRegex(ValueError, 'PLAN_BINDING_MISMATCH'): w.prepare(args)

class BrowserGuards(unittest.TestCase):
    def setUp(self):
        self.row = {'key': 'basic.name', 'selector': '#name', 'label': 'Name', 'kind': 'text', 'value_from': 'profile:basic.name', 'expected': ''}
        self.meta = {'tag': 'input', 'type': 'text', 'name': 'name', 'label': 'Name', 'maxLength': None, 'disabled': False, 'readonly': False, 'value': ''}

    def check(self, actual='', value='Alex Example'):
        with patch.object(b, 'rf_inspect', return_value=self.meta), patch.object(b, 'rf_read', return_value=actual):
            return b.rf_preflight(self.row, value)

    def test_nonempty_conflict(self):
        self.row['expected'] = 'User edit'
        with self.assertRaisesRegex(RuntimeError, 'NONEMPTY_CONFLICT'): self.check('User edit')

    def test_page_changed(self):
        with self.assertRaisesRegex(RuntimeError, 'PAGE_VALUE_CHANGED'): self.check('User edit')

    def test_idempotent_repeat(self):
        self.assertEqual(self.check('Alex Example'), 'already_equal')

    def test_password_and_file_inputs(self):
        for kind in ('password', 'file'):
            self.meta['type'] = kind
            with self.assertRaises(RuntimeError): self.check()

    def test_consent_is_manual(self):
        self.meta['label'] = 'Name privacy consent'
        with self.assertRaisesRegex(RuntimeError, 'CONSENT_REQUIRES_USER'): self.check()

    def test_maxlength_does_not_truncate(self):
        self.meta['maxLength'] = 3
        with self.assertRaisesRegex(RuntimeError, 'TOO_LONG'): self.check()

    def test_readonly_does_not_mutate(self):
        self.meta['readonly'] = True
        with self.assertRaisesRegex(RuntimeError, 'DISABLED_OR_READONLY'): self.check()

if __name__ == '__main__':
    unittest.main()
